"""Перевод документов: TXT, SRT, VTT, DOCX, PDF."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable

from ..domain.errors import ServiceError
from ..localization import tr
from .translation import TranslationService

Progress = Callable[[int, int, str], None]
Cancel = Callable[[], bool]

SUPPORTED_EXTENSIONS = (".txt", ".srt", ".vtt", ".docx", ".pdf")
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1251")


class DocumentCancelled(Exception):
    """Перевод документа прерван пользователем."""


def supported_extensions() -> list[str]:
    return list(SUPPORTED_EXTENSIONS)


def output_extension(extension: str) -> str:
    """PDF отдаём текстом, остальные форматы сохраняют своё расширение."""
    return ".txt" if extension.lower() == ".pdf" else extension


class DocumentTranslator:
    def __init__(self, service: TranslationService, provider_id: str) -> None:
        self._service = service
        self._provider_id = provider_id

    def translate_file(
        self,
        input_path: Path,
        output_path: Path,
        source: str,
        target: str,
        progress: Progress,
        cancel: Cancel,
    ) -> None:
        extension = input_path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ServiceError(
                tr("errors.unsupported_format", extension=extension or "?"), kind="unsupported"
            )
        if not input_path.is_file():
            raise ServiceError(tr("errors.file_not_found"), kind="config")

        def translate(text: str) -> str:
            return self._translate(text, source, target)

        if extension == ".txt":
            output_path.write_text(
                self._text(input_path, translate, progress, cancel), encoding="utf-8"
            )
        elif extension in (".srt", ".vtt"):
            output_path.write_text(
                self._subtitles(input_path, translate, progress, cancel), encoding="utf-8"
            )
        elif extension == ".docx":
            self._docx(input_path, output_path, translate, progress, cancel)
        elif extension == ".pdf":
            self._pdf(input_path, output_path, translate, progress, cancel)


    def _text(self, path: Path, translate, progress: Progress, cancel: Cancel) -> str:
        paragraphs = re.split(r"\n\s*\n", self._read_text(path))
        total = len(paragraphs)
        translated: list[str] = []
        for index, paragraph in enumerate(paragraphs, start=1):
            if cancel():
                raise DocumentCancelled()
            translated.append(translate(paragraph) if paragraph.strip() else paragraph)
            progress(index, total, "текст")
        return "\n\n".join(translated)

    def _subtitles(self, path: Path, translate, progress: Progress, cancel: Cancel) -> str:
        lines = self._read_text(path).replace("\r\n", "\n").split("\n")
        output = list(lines)

        ranges: list[tuple[int, int]] = []
        index = 0
        while index < len(lines):
            if "-->" in lines[index]:
                end = index + 1
                while end < len(lines) and lines[end].strip() != "":
                    end += 1
                ranges.append((index + 1, end))
                index = end
            else:
                index += 1

        total = len(ranges)
        for position, (start, end) in enumerate(ranges, start=1):
            if cancel():
                raise DocumentCancelled()
            content = "\n".join(lines[start:end]).strip()
            if content:
                output[start:end] = translate(content).split("\n")
            progress(position, total, "субтитры")
        return "\n".join(output)

    def _docx(self, path: Path, output_path: Path, translate, progress: Progress, cancel: Cancel) -> None:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise ServiceError(tr("errors.docx_missing"), kind="config") from exc

        document = docx.Document(str(path))
        paragraphs = list(self._docx_paragraphs(document))
        total = len(paragraphs)
        for index, paragraph in enumerate(paragraphs, start=1):
            if cancel():
                raise DocumentCancelled()
            text = paragraph.text
            if text.strip():
                self._set_paragraph_text(paragraph, translate(text))
            progress(index, total, "docx")
        document.save(str(output_path))

    def _pdf(self, path: Path, output_path: Path, translate, progress: Progress, cancel: Cancel) -> None:
        try:
            import pypdf
        except ImportError as exc:  # pragma: no cover
            raise ServiceError(tr("errors.pdf_missing"), kind="config") from exc

        reader = pypdf.PdfReader(str(path))
        total = len(reader.pages)
        chunks: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            if cancel():
                raise DocumentCancelled()
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(translate(text))
            progress(index, total, "pdf")
        output_path.write_text("\n\n".join(chunks), encoding="utf-8")


    def _translate(self, text: str, source: str, target: str) -> str:
        results = self._service.translate(text, source, target, [self._provider_id])
        best = next((result for result in results if result.ok), None)
        if best is None:
            message = results[0].error if results else tr("errors.no_translator")
            raise ServiceError(message or tr("errors.translation_failed"), kind="error")
        return best.text

    @staticmethod
    def _read_text(path: Path) -> str:
        for encoding in _TEXT_ENCODINGS:
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ServiceError(tr("errors.file_encoding"), kind="config")

    @staticmethod
    def _docx_paragraphs(document):
        for paragraph in document.paragraphs:
            yield paragraph
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph

    @staticmethod
    def _set_paragraph_text(paragraph, text: str) -> None:
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(text)
