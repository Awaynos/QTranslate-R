"""Проверка перевода документов: TXT, SRT, VTT, DOCX."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from qtranslate_r.core.documents import DocumentTranslator  # noqa: E402
from qtranslate_r.core.setup import build_registry  # noqa: E402
from qtranslate_r.core.translation import TranslationService  # noqa: E402

WORK = Path(tempfile.mkdtemp(prefix="qtr_docs_"))
PROVIDER = "google"


def translate_file(source: Path, target: Path) -> str:
    translator = DocumentTranslator(TranslationService(build_registry()), PROVIDER)
    translator.translate_file(source, target, "en", "ru", lambda *a: None, lambda: False)
    return target.read_text(encoding="utf-8")


def test_txt() -> bool:
    source = WORK / "sample.txt"
    source.write_text("Hello world.\n\nThis is a test document.", encoding="utf-8")
    result = translate_file(source, WORK / "sample.out.txt")
    print("TXT:", result.replace("\n", " | "))
    return "Привет" in result or "мир" in result.lower()


def test_srt() -> bool:
    source = WORK / "sample.srt"
    source.write_text(
        "1\n00:00:01,000 --> 00:00:03,000\nHello world.\n\n"
        "2\n00:00:04,000 --> 00:00:06,000\nGood morning.",
        encoding="utf-8",
    )
    result = translate_file(source, WORK / "sample.out.srt")
    print("SRT:", result.replace("\n", " | ")[:160])
    return "00:00:01,000 --> 00:00:03,000" in result and "Hello world." not in result


def test_vtt() -> bool:
    source = WORK / "sample.vtt"
    source.write_text(
        "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nHello world.\n",
        encoding="utf-8",
    )
    result = translate_file(source, WORK / "sample.out.vtt")
    print("VTT:", result.replace("\n", " | ")[:160])
    return result.startswith("WEBVTT") and "Hello world." not in result


def test_docx() -> bool:
    try:
        import docx
    except ImportError:
        print("DOCX: python-docx не установлен — пропуск")
        return True

    source = WORK / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Hello world.")
    document.add_paragraph("This is a test document.")
    document.save(str(source))

    target = WORK / "sample.out.docx"
    translator = DocumentTranslator(TranslationService(build_registry()), PROVIDER)
    translator.translate_file(source, target, "en", "ru", lambda *a: None, lambda: False)

    result = docx.Document(str(target))
    texts = [p.text for p in result.paragraphs if p.text.strip()]
    print("DOCX:", texts)
    return bool(texts) and all(t != "Hello world." for t in texts)


def main() -> int:
    checks = {
        "txt": test_txt(),
        "srt": test_srt(),
        "vtt": test_vtt(),
        "docx": test_docx(),
    }
    print("Итог:", checks)
    ok = all(checks.values())
    print("OK" if ok else "ПРОВАЛ")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
